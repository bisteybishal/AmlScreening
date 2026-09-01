from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("AML_Screening_Portal_Project_Documentation.docx")

# compact_reference_guide preset, plus a named AML teal accent override used
# only for the cover kicker and callouts.
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "0E918C"
PALE_TEAL = "E7F5F3"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK_GRAY = "344054"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_run_font(run, name="Calibri", size=11, color="000000", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must total {PAGE_WIDTH_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def set_paragraph_border(paragraph, color=BLUE, size=10, space=4, side="left"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    borders.append(border)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "\uf0b7" if kind == "bullet" else "%1.")
    lvl.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Symbol")
        fonts.set(qn("w:hAnsi"), "Symbol")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    # Word requires every abstractNum to appear before the first concrete num.
    # Appending abstractNum after the template's existing num elements causes
    # Word to repair the file and can remap bullet lists to decimal numbering.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle.paragraph_format.space_after = Pt(18)

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("AML Screening Portal  |  Project Documentation")
    set_run_font(run, size=8.5, color=MID_GRAY, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)


def add_paragraph(doc, text="", *, bold_prefix=None, italic=False, color="000000", after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=color)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, italic=italic, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic, color=color)
    return p


def add_bullet(doc, text, bullet_num_id):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_numbered(doc, text, number_num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    apply_numbering(p, number_num_id)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_callout(doc, label, text, fill=PALE_TEAL, border=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = True
    shade_paragraph(p, fill)
    set_paragraph_border(p, border, size=18, space=8)
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=NAVY)
    text_run = p.add_run(text)
    set_run_font(text_run, color=DARK_GRAY)


def add_code_block(doc, code):
    lines = code.strip("\n").splitlines()
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(0 if index else 5)
        p.paragraph_format.space_after = Pt(0 if index < len(lines) - 1 else 8)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = index < len(lines) - 1
        shade_paragraph(p, LIGHT_GRAY)
        if index == 0:
            set_paragraph_border(p, BLUE, size=12, space=6)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=8.3, color=NAVY)


def add_table(doc, headers, rows, widths_dxa, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=DARK_GRAY)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    return table


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=NAVY)
    value_run = p.add_run(value)
    set_run_font(value_run, color=DARK_GRAY)


def add_menu_entry(doc, name, route, permission, purpose, actions, tables, bullet_num_id, note=None):
    doc.add_heading(name, level=2)
    add_label_value(doc, "Route", route)
    add_label_value(doc, "Required permission", permission)
    add_paragraph(doc, purpose, color=DARK_GRAY)
    add_label_value(doc, "Main operations", "")
    for action in actions:
        add_bullet(doc, action, bullet_num_id)
    add_label_value(doc, "Primary tables", tables)
    if note:
        add_callout(doc, "Important", note, fill="FFF7E6", border=GOLD)


def build_document():
    doc = Document()
    configure_document(doc)
    bullet_num_id = add_numbering_definition(doc, "bullet")
    contents_num_id = add_numbering_definition(doc, "decimal")
    login_num_id = add_numbering_definition(doc, "decimal")
    workflow_num_id = add_numbering_definition(doc, "decimal")
    migration_num_id = add_numbering_definition(doc, "decimal")
    priority_num_id = add_numbering_definition(doc, "decimal")

    # Cover page - editorial_cover pattern.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AML SCREENING")
    set_run_font(r, size=13, color=TEAL, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Portal Project Documentation")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Architecture, Menus, Database, Stored Procedures and Operational Flows")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Implementation reference for Development, UAT and Production preparation")
    set_run_font(r, size=11, color=DARK_GRAY, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(88)
    for label, value in (
        ("Document version", "1.0"),
        ("Prepared", "1 September 2026"),
        ("Application", "AmlScreening.Portal (.NET 10 MVC)"),
        ("Database", "SQL Server - AMLScreening"),
        ("Scope", "Portal implementation; client API remains a separate future project"),
    ):
        row = doc.add_paragraph()
        row.alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.paragraph_format.space_after = Pt(3)
        rr = row.add_run(f"{label}: ")
        set_run_font(rr, size=10, color=MID_GRAY, bold=True)
        rr = row.add_run(value)
        set_run_font(rr, size=10, color=DARK_GRAY)
    doc.add_page_break()

    doc.add_heading("Document control", level=1)
    add_table(doc,
              ["Item", "Detail"],
              [
                  ("Purpose", "Describe the implemented AML Screening portal for developers, testers, administrators and compliance users."),
                  ("System status", "Portal workflow and database migration are implemented. The external client API is explicitly deferred."),
                  ("Evidence basis", "Current source code, applied SQL Server schema, live menu/permission catalog and passing integration tests."),
                  ("Confidentiality", "Internal project and controlled client/UAT use."),
              ],
              [2700, 6660], font_size=9.5)

    add_callout(doc, "Current state",
                "The live database contains 19 dbo tables and 3 stored procedures. The portal uses EF Core and Dapper inline queries for its active paths; the procedures are documented as legacy/support objects and are not currently invoked by the C# repositories.")

    doc.add_heading("Contents", level=1)
    contents = [
        "Project purpose and scope",
        "Solution architecture and responsibilities",
        "Authentication, authorization and tenant isolation",
        "Menu-by-menu functional guide",
        "End-to-end AML and compliance workflow",
        "Risk scoring and decision logic",
        "Database architecture and table dictionary",
        "Stored procedure catalog and SQL definitions",
        "Data access, audit, soft delete and concurrency",
        "Permissions and access-control matrix",
        "Configuration, migration, UAT and production guidance",
        "Testing evidence, limitations and remaining work",
        "Appendices: project structure, scripts and glossary",
    ]
    for item in contents:
        add_numbered(doc, item, contents_num_id)

    doc.add_heading("1. Project purpose and scope", level=1)
    add_paragraph(doc,
                  "The AML Screening Portal is a multi-client web application for managing screening lists, customer KYC profiles, compliance transactions, risk rules, screening evidence, automated decisions, alerts, cases, users, roles and audit history. It is designed as the operational portal that compliance staff use before a separate integration API is introduced.")
    add_paragraph(doc,
                  "The implemented prototype supports a practical transaction journey from customer creation to ALLOW, REVIEW or BLOCK. It is suitable for development and controlled UAT after environment configuration, but it is not yet a complete production sanctions platform because automated list ingestion, maker-checker approval, client API credentials and advanced transaction-monitoring models remain future work.")

    doc.add_heading("Implemented scope", level=2)
    implemented = [
        "Cookie-based login with username/IP rate limiting and password hashing.",
        "Role permissions plus per-user grant or deny overrides.",
        "Database-driven sidebar menu definitions.",
        "Client-level query isolation and AML-record ownership controls.",
        "AML record add, edit, pagination and soft delete.",
        "Customer/KYC management and initial customer risk scoring.",
        "Transaction capture, screening, rule evaluation and persisted decisions.",
        "Compliance alerts, cases and audit records.",
        "Subscription plan management with soft delete and concurrency handling.",
        "Separate Development, UAT and Production configuration placeholders.",
        "SQL Server and in-memory integration tests.",
    ]
    for item in implemented:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("Deferred scope", level=2)
    deferred = [
        "External client API project, API credentials, OAuth2/mTLS and idempotency.",
        "Automated PEP, UN, EU, CIB and sanctions-list import/version pipelines.",
        "Country-risk masters, customer baselines, structuring detection and network analytics.",
        "Maker-checker approval workflow for rule and watchlist changes.",
        "Evidence attachments, regulatory STR/SAR submission and production case notes timeline.",
        "Distributed cache, message bus, outbox workers and search engine integration at scale.",
    ]
    for item in deferred:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("2. Solution architecture and responsibilities", level=1)
    add_code_block(doc, """
Browser / Razor Views
        -> MVC Controllers and authorization policies
        -> Application interfaces and authentication models
        -> Infrastructure repositories and workflow services
        -> EF Core / Dapper / SQL Server
        -> Audit evidence and client-scoped records
""")
    architecture_rows = [
        ("AmlScreening.Portal", "MVC controllers, Razor views, authentication cookie, rate limiter, authorization handlers and current-user context."),
        ("AmlScreening.Application", "Interfaces, permission constants, login service, DTOs and workflow result models."),
        ("AmlScreening.Domain", "Business entities for users, clients, AML records, customers, transactions, rules, decisions, alerts and cases."),
        ("AmlScreening.Infrastructure", "EF Core DbContext, Dapper connection factory, repositories, menu service, audit service and AML workflow engine."),
        ("Database", "Versioned SQL scripts for isolation, permissions, menus, compliance rules, workflow tables and deployment preparation."),
        ("IntegrationTests", "Cross-client checks, schema checks and repository behavior tests."),
        ("DatabaseMigrator", "Small command-line runner for controlled execution of SQL migration scripts."),
    ]
    add_table(doc, ["Layer/project", "Responsibility"], architecture_rows, [2700, 6660], font_size=9.4)

    doc.add_heading("Design principles used", level=2)
    principles = [
        "Single responsibility: controllers coordinate HTTP behavior; repositories/services perform data and business work.",
        "Dependency inversion: controllers depend on application interfaces rather than concrete SQL implementations.",
        "Encapsulation: current user, client and permissions are exposed through ICurrentUserContext.",
        "Open/closed direction: new menu items, permissions and rules can be added without rewriting the layout.",
        "Defense in depth: menu filtering improves usability while controller policies and query scoping enforce security.",
    ]
    for item in principles:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("3. Authentication, authorization and tenant isolation", level=1)
    doc.add_heading("Login flow", level=2)
    for step in [
        "The user submits a username, password and optional remember-me selection.",
        "The login rate limiter partitions attempts by normalized username and remote IP address.",
        "UserRepository loads the active user and active role from Users, UserRoles and Roles.",
        "PasswordService verifies the stored password hash in application code.",
        "Role permissions are combined with UserPermissions overrides. IsGranted allows an explicit grant or deny.",
        "The portal creates claims for user ID, ClientId, username, full name, role and permissions.",
        "The secure authentication cookie is issued and the user is redirected to Dashboard.",
        "A login audit event is written; an audit failure is logged without preventing a valid login.",
    ]:
        add_numbered(doc, step, login_num_id)

    add_callout(doc, "Rate limit",
                "The current fixed window permits five login POST attempts per username and IP combination per minute. Browser tabs share the same cookie, so opening more tabs does not create separate login sessions.")

    doc.add_heading("Permission resolution", level=2)
    add_code_block(doc, """
Role permission exists and no user override -> granted
Role permission exists and user override IsGranted = 0 -> denied
Role permission absent and user override IsGranted = 1 -> granted
Neither role permission nor user grant exists -> denied
""")
    add_paragraph(doc,
                  "Menu visibility is based on MenuItems.PermissionCode. Actual access is separately enforced by [Authorize(Policy = ...)] on controllers and actions. A hidden menu is therefore not treated as the security boundary.")

    doc.add_heading("Client and owner scope", level=2)
    scope_rows = [
        ("ClientId", "Separates organizations. Client-scoped data queries are limited to the logged-in user's ClientId unless the platform-wide permission applies."),
        ("OwnerUserId", "Used by AmlRecords to restrict ordinary users to records they created."),
        ("AMLRecords.ViewClient", "Allows a user to view all AML records belonging to the current client, including records created by an administrator."),
        ("AMLRecords.ViewAll", "Platform-level access across clients. This permission also bypasses EF Core client filters and must be assigned only to trusted platform administrators."),
    ]
    add_table(doc, ["Control", "Use"], scope_rows, [2600, 6760], font_size=9.4)
    add_callout(doc, "Current limitation",
                "Customers, transactions, alerts and cases are client-wide rather than owner-user-specific. If each client employee must see only their own customer workflow records, add OwnerUserId and resource authorization to those tables and services.", fill="FFF7E6", border=GOLD)

    doc.add_heading("4. Menu-by-menu functional guide", level=1)
    add_paragraph(doc,
                  "Menu definitions are loaded from dbo.MenuItems. ParentMenuId creates the Compliance submenu. PermissionCode controls visibility, while ControllerName and ActionName identify the MVC destination.")

    add_menu_entry(doc, "Dashboard", "/Dashboard/Index", "Authenticated user",
                   "Provides a summary landing page after login. It combines AML record metrics and compliance workflow counts.",
                   ["Show total and active AML records.", "Show transaction count and decisions requiring review or block.", "Show open alert and open case counts.", "Display identity and role information from authentication claims."],
                   "AmlRecords, AmlTransactions, ComplianceAlerts, ComplianceCases", bullet_num_id,
                   "Dashboard values are now loaded from repositories/services rather than remaining fully hardcoded. Trend visuals may still require historical aggregation for production accuracy.")

    add_menu_entry(doc, "AML Records", "/AmlRecords/Index", "AMLRecords.View",
                   "Maintains the AML screening list data used during transaction evaluation. Records can represent individuals or entities from PEP, UN, EU, CIB and other sources.",
                   ["View paginated records.", "Search by name, reference or source.", "Filter by record type and active status.", "Add records with AMLRecords.Create.", "Edit records with AMLRecords.Edit.", "Soft delete records with AMLRecords.Delete."],
                   "AmlRecords, AuditLogs", bullet_num_id,
                   "Ordinary users see their own records unless AMLRecords.ViewClient is granted. Client-wide master data should normally be managed by trusted users and shared within the same client.")

    add_menu_entry(doc, "Users", "/Users/Index", "Users.View",
                   "Creates and lists portal users for the permitted client scope. Clicking the user name opens individual access editing.",
                   ["List active and inactive users with pagination.", "Create a user and assign a role with Users.Create.", "Open individual access from the user name.", "Apply per-user permission grants or denies with Roles.Manage."],
                   "Users, UserRoles, Roles, UserPermissions, Permissions, AuditLogs", bullet_num_id)

    add_menu_entry(doc, "Roles & Permissions", "/Roles/Index", "Roles.Manage",
                   "Centralizes reusable access profiles such as Admin, AML User and Viewer.",
                   ["Create a role.", "View all available permissions.", "Assign or remove permissions from a role.", "Use role permissions as the default access baseline for users."],
                   "Roles, Permissions, RolePermissions, AuditLogs", bullet_num_id,
                   "Role permissions should hold common access. UserPermissions should be reserved for exceptions, which keeps access administration understandable.")

    add_menu_entry(doc, "Subscription", "/Subscription/Index", "Subscription.View",
                   "Maintains source subscription plans, durations and rates. This is the portal-side plan catalog for later API entitlement work.",
                   ["List and paginate plans.", "Create a plan with Subscription.Create.", "Edit a plan with Subscription.Edit.", "Soft delete a plan with Subscription.Delete.", "Use RowVersion to prevent silent concurrent overwrites."],
                   "AML_Subscription_Plan, AuditLogs", bullet_num_id,
                   "Subscription data is not yet connected to API credentials, usage quotas or transaction enforcement. Those belong to the future API project.")

    doc.add_heading("Compliance", level=2)
    add_paragraph(doc,
                  "Compliance is a parent menu without its own controller. Its children form the operational transaction-monitoring workflow.")

    add_menu_entry(doc, "Compliance > Customers", "/Customers/Index", "Customers.View",
                   "Maintains customer/KYC profiles and their initial risk indicators.",
                   ["View client customers.", "Create customer identity and KYC details.", "Calculate customer risk from PEP status and missing KYC factors.", "Edit with optimistic concurrency.", "Soft delete customers."],
                   "Customers, Clients, AuditLogs", bullet_num_id)

    add_menu_entry(doc, "Compliance > Transactions", "/Transactions/Index", "Transactions.View",
                   "Captures customer transactions and provides the Evaluate action that runs screening, rules, scoring and decisioning.",
                   ["List recent client transactions.", "Create a transaction linked to an existing customer.", "View transaction, customer, match, risk and decision details.", "Evaluate with Transactions.Evaluate.", "Reuse an existing persisted result when a transaction has already been evaluated."],
                   "AmlTransactions, Customers, ScreeningMatches, RiskAssessments, AmlDecisions, ComplianceAlerts, ComplianceCases, AuditLogs", bullet_num_id,
                   "The current list is capped at 500 recent records and should receive server-side pagination before large production volumes.")

    add_menu_entry(doc, "Compliance > Threshold Rules", "/ComplianceRules/Index", "ComplianceRules.View",
                   "Configures the transaction rules used by the AML workflow engine.",
                   ["Create amount-threshold or velocity rules.", "Set currency, threshold, lookback, transaction count and risk score.", "Set the resulting action: Allow, Review/Hold/Alert or Block.", "Control priority, rule version, approval status and effective dates.", "Soft delete obsolete rules."],
                   "ComplianceRules, Clients, AuditLogs", bullet_num_id,
                   "ApprovalStatus is currently editable directly; true maker-checker approval is recommended before production.")

    add_menu_entry(doc, "Compliance > Alerts", "/Alerts/Index", "Alerts.View",
                   "Shows generated compliance alerts that require analyst attention.",
                   ["View severity, reason, transaction and status.", "Assign an alert to a user.", "Move the alert through investigation statuses.", "Require a resolution when resolving or closing an alert with Alerts.Manage."],
                   "ComplianceAlerts, AmlTransactions, RiskAssessments, Users, AuditLogs", bullet_num_id)

    add_menu_entry(doc, "Compliance > Cases", "/Cases/Index", "Cases.View",
                   "Provides the formal investigation record created from an alert.",
                   ["View cases and detailed evidence.", "Assign an investigator.", "Record status, disposition and notes.", "Require a disposition for cleared, rejected, STR/SAR or closed outcomes with Cases.Manage."],
                   "ComplianceCases, ComplianceAlerts, AmlTransactions, Users, AuditLogs", bullet_num_id)

    add_menu_entry(doc, "Audit Logs", "/AuditLogs/Index", "AuditLogs.View",
                   "Provides traceability for important portal and workflow actions.",
                   ["View recent login, logout, create, update, delete, evaluate, alert and case actions.", "Restrict ordinary viewers to their client.", "Allow controlled platform-wide review for appropriately privileged administrators."],
                   "AuditLogs, Clients, Users", bullet_num_id)

    doc.add_heading("Removed or inactive menu concepts", level=2)
    add_paragraph(doc,
                  "Search Screening, Reports and Settings are not present in the active MenuItems catalog. Their permission names remain in dbo.Permissions for backward compatibility but should be retired or reintroduced only when working controllers and views exist.")

    doc.add_heading("5. End-to-end AML and compliance workflow", level=1)
    add_code_block(doc, """
AML list data -> Customer/KYC -> Transaction -> Screening matches
              -> Compliance rules -> Risk assessment -> Decision
              -> ALLOW, or REVIEW/BLOCK -> Alert -> Case -> Resolution
""")
    workflow_steps = [
        ("Prepare screening data", "An authorized user adds active PEP, UN, EU, CIB or other watchlist entries to AmlRecords."),
        ("Create customer", "A customer reference and KYC details are saved. Customer risk is calculated immediately."),
        ("Create transaction", "A transaction is linked to the customer, assigned a correlation ID and stored with status NEW."),
        ("Screen parties", "Sender and receiver names are normalized and compared with active AmlRecords for the current client."),
        ("Apply rules", "Active and effective AmountThreshold and Velocity rules are evaluated in priority order."),
        ("Calculate risk", "Customer, transaction and screening scores are combined and persisted in RiskAssessments."),
        ("Decide", "AmlDecisions records ALLOW, REVIEW or BLOCK together with reasons, model version, rule versions and list versions."),
        ("Escalate", "REVIEW and BLOCK create a ComplianceAlert and ComplianceCase for analyst work."),
        ("Investigate and close", "Compliance staff assign, investigate, document and resolve the alert/case."),
        ("Audit", "Application actions are written to AuditLogs throughout the journey."),
    ]
    for title_text, detail in workflow_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        apply_numbering(p, workflow_num_id)
        r = p.add_run(f"{title_text}. ")
        set_run_font(r, bold=True, color=NAVY)
        r = p.add_run(detail)
        set_run_font(r, color=DARK_GRAY)

    doc.add_heading("Example transaction", level=2)
    example_rows = [
        ("AML record", "UN-000125, John Example, Individual, UN Sanctions"),
        ("Customer", "CUS-1001, John Example, PEP = Yes"),
        ("Transaction", "TXN-9001, NPR 1,200,000, Nepal to destination country"),
        ("Screening", "Name similarity 88%; potential match stored"),
        ("Rule", "AmountThreshold rule triggers Review and adds transaction risk"),
        ("Result", "Risk evidence persisted; REVIEW decision; alert and case opened"),
    ]
    add_table(doc, ["Stage", "Example result"], example_rows, [2200, 7160], font_size=9.4)

    doc.add_page_break()
    doc.add_heading("6. Risk scoring and decision logic", level=1)
    add_callout(doc, "Prototype formula",
                "Overall risk = Customer risk x 40% + Transaction risk x 40% + Screening risk x 20%.")

    risk_rows = [
        ("Customer risk", "PEP adds 50; missing source of funds adds 15; missing occupation adds 10; final score is capped at 100."),
        ("Screening risk", "Levenshtein name similarity. Sender match can gain 10 for matching date of birth and 5 for matching nationality. Below 75 is ignored; 75-94.99 is potential; 95-100 is confirmed."),
        ("Transaction risk", "Triggered rules add configured RiskScore. NPR/amount-independent fallback: amount >= 1,000,000 sets at least 60; amount >= 500,000 sets at least 35."),
        ("Risk level", "0-29 LOW; 30-59 MEDIUM; 60-79 HIGH; 80-100 CRITICAL."),
        ("Decision", "Confirmed match or Block action -> BLOCK. Potential match, score >= 60, or Review/Hold/Alert action -> REVIEW. Otherwise -> ALLOW."),
    ]
    add_table(doc, ["Component", "Current implementation"], risk_rows, [2200, 7160], font_size=9.2)
    add_callout(doc, "Production caution",
                "This score is an explainable prototype, not a regulator-approved model. Before production it requires governance, validated thresholds, country/product/channel risk, false-positive testing, model approval and periodic review.", fill="FDECEC", border=RED)

    doc.add_heading("7. Database architecture and table dictionary", level=1)
    add_paragraph(doc,
                  "The applied SQL Server database contains 19 dbo tables. ClientId is the principal tenant boundary. Foreign keys preserve core relationships, filtered unique indexes protect active references, DeleteFg supports soft deletion and RowVersion supports optimistic concurrency where configured.")

    doc.add_heading("Core relationships", level=2)
    add_code_block(doc, """
Clients 1 -> many Users, AmlRecords, Customers, Transactions, Rules, Alerts, Cases
Users many -> many Roles through UserRoles
Roles many -> many Permissions through RolePermissions
Users many -> many permission overrides through UserPermissions
Customers 1 -> many AmlTransactions
AmlTransactions 1 -> many ScreeningMatches
AmlTransactions 1 -> 1 RiskAssessment -> 1 AmlDecision
AmlTransactions -> ComplianceAlert -> ComplianceCase
MenuItems parent -> child MenuItems
""")

    doc.add_heading("Identity, access and navigation tables", level=2)
    identity_tables = [
        ("Clients", "Tenant/organization master.", "Id PK, ClientCode unique, ClientName, IsActive, CreatedDate. Parent of client-owned business data."),
        ("Users", "Portal login accounts.", "UserId PK, ClientId FK, Username, FullName, PasswordHash, IsActive, LastLoginDate, audit dates."),
        ("Roles", "Reusable access profiles.", "RoleId PK, RoleName, IsActive, CreatedDate."),
        ("UserRoles", "User-to-role bridge.", "Composite relationship: UserId FK -> Users; RoleId FK -> Roles."),
        ("Permissions", "Named application capabilities.", "PermissionId PK, PermissionName unique, Description, IsActive."),
        ("RolePermissions", "Role permission assignments.", "Composite PK; RoleId FK -> Roles; PermissionId FK -> Permissions."),
        ("UserPermissions", "Per-user grant/deny overrides.", "Composite PK; UserId FK, PermissionId FK, IsGranted. Overrides the role result."),
        ("MenuItems", "Database-driven sidebar catalog.", "MenuId PK, ParentMenuId self-FK, MenuKey, DisplayName, controller/action, icon, PermissionCode, order, active flag."),
    ]
    add_table(doc, ["Table", "Purpose", "Important columns and relationships"], identity_tables, [1900, 2500, 4960], font_size=8.8)

    doc.add_heading("AML and compliance workflow tables", level=2)
    workflow_tables = [
        ("AmlRecords", "PEP/sanctions/watchlist master data.", "Id PK, ClientId FK, OwnerUserId FK, type/name/DOB/country/nationality/reference/source, active/audit/soft-delete/concurrency fields."),
        ("Customers", "Customer/KYC profiles and customer risk.", "Id PK, ClientId FK, CustomerReference unique per client, KYC fields, IsPep, RiskScore, RiskLevel, lifecycle fields."),
        ("AmlTransactions", "Transactions submitted for evaluation.", "Id PK, ClientId FK, CustomerId FK, reference, parties, amount/currency, countries/channel/date, correlation, decision and lifecycle fields."),
        ("ComplianceRules", "Tenant rule configuration.", "Id PK, ClientId FK, code/version/type/threshold/velocity/action/risk/priority/approval/effective dates, lifecycle fields."),
        ("ScreeningMatches", "Evidence of watchlist comparisons.", "Id PK, ClientId FK, TransactionId FK, AmlRecordId FK, subject/matched names, score/status/source/list version/fields."),
        ("RiskAssessments", "Persisted component and overall risk.", "Id PK, ClientId FK, TransactionId unique FK, component scores, overall score/level, model version, factors and calculation time."),
        ("AmlDecisions", "Explainable ALLOW/REVIEW/BLOCK result.", "Id PK, ClientId FK, TransactionId unique FK, RiskAssessmentId FK, correlation, reasons, rule/model/list versions."),
        ("ComplianceAlerts", "Analyst alert work queue.", "Id PK, ClientId FK, TransactionId/RiskAssessmentId FKs, severity/reason/status, AssignedUserId FK, resolution and lifecycle fields."),
        ("ComplianceCases", "Formal investigation record.", "Id PK, ClientId FK, AlertId/TransactionId FKs, status, AssignedUserId FK, disposition, notes, closed date and lifecycle fields."),
    ]
    add_table(doc, ["Table", "Purpose", "Important columns and relationships"], workflow_tables, [1900, 2500, 4960], font_size=8.6)

    doc.add_page_break()
    doc.add_heading("Operational and commercial tables", level=2)
    operational_tables = [
        ("AuditLogs", "Immutable-style application activity evidence.", "AuditLogId PK, ClientId FK, UserId, Username, Action, EntityType, EntityId, Details, CreatedDate."),
        ("AML_Subscription_Plan", "Subscription/source plan catalog.", "PlanId PK, SourceCode/SourceName, DurationMonths, Rate, IsActive, audit/soft-delete/concurrency fields."),
    ]
    add_table(doc, ["Table", "Purpose", "Important columns and relationships"], operational_tables, [1900, 2500, 4960], font_size=8.9)

    doc.add_heading("Data integrity controls", level=2)
    for item in [
        "Unique client references prevent duplicate active AML, customer, transaction, alert and case references within the same client.",
        "Foreign keys prevent orphaned client, customer, transaction, alert, user and permission relationships.",
        "Check constraints enforce positive transaction amounts, valid risk-score ranges and Y/N soft-delete flags.",
        "RowVersion detects concurrent edits for supported entities.",
        "EF Core query filters hide soft-deleted rows and enforce ClientId scope for mapped workflow entities.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("8. Stored procedure catalog and SQL definitions", level=1)
    add_callout(doc, "Verified live catalog",
                "Three stored procedures currently exist. Source-code search found no current C# call using CommandType.StoredProcedure or these procedure names. Active authentication and repositories use inline Dapper SQL or EF Core.")
    procedure_rows = [
        ("spa_AML_SubscriptionPlan_GetAll", "Returns active subscription plans ordered by PlanId.", "Legacy/support; not called by current SubscriptionRepository."),
        ("usp_User_GetForLogin", "Returns an active user and role by username.", "Legacy; missing ClientId and permission resolution required by current login result."),
        ("usp_User_UpdateLastLogin", "Updates LastLoginDate for a user.", "Legacy; current repository performs an inline SYSUTCDATETIME update."),
    ]
    add_table(doc, ["Procedure", "Purpose", "Current status"], procedure_rows, [2900, 3100, 3360], font_size=8.8)

    doc.add_heading("dbo.spa_AML_SubscriptionPlan_GetAll", level=2)
    add_code_block(doc, """
CREATE PROC dbo.spa_AML_SubscriptionPlan_GetAll
AS
BEGIN
    SET NOCOUNT ON;
    SELECT PlanId, SourceCode, SourceName, DurationMonths, Rate, IsActive
    FROM dbo.AML_Subscription_Plan
    WHERE IsActive = 1
    ORDER BY PlanId;
END;
""")
    add_paragraph(doc,
                  "Before reuse, add DeleteFg filtering and decide whether subscription plans are global or client-owned. Current repository behavior includes pagination, soft delete, auditing and concurrency that this procedure does not provide.")

    doc.add_heading("dbo.usp_User_GetForLogin", level=2)
    add_code_block(doc, """
CREATE PROCEDURE dbo.usp_User_GetForLogin
    @Username VARCHAR(100)
AS
BEGIN
    SET NOCOUNT ON;
    SELECT u.UserId, u.Username, u.FullName, u.PasswordHash,
           u.IsActive, r.RoleId, r.RoleName
    FROM dbo.Users u
    INNER JOIN dbo.UserRoles ur ON u.UserId = ur.UserId
    INNER JOIN dbo.Roles r ON ur.RoleId = r.RoleId
    WHERE u.Username = @Username
      AND u.IsActive = 1
      AND r.IsActive = 1;
END;
""")
    add_callout(doc, "Do not switch directly",
                "The current login requires ClientId and separately resolves role permissions plus UserPermissions.IsGranted overrides. This stored procedure must be updated and regression-tested before it can replace UserRepository.GetUserAsync.", fill="FFF7E6", border=GOLD)

    doc.add_heading("dbo.usp_User_UpdateLastLogin", level=2)
    add_code_block(doc, """
CREATE PROCEDURE dbo.usp_User_UpdateLastLogin
    @UserId BIGINT
AS
BEGIN
    SET NOCOUNT ON;
    UPDATE dbo.Users
    SET LastLoginDate = GETDATE()
    WHERE UserId = @UserId;
END;
""")
    add_paragraph(doc,
                  "The active repository uses SYSUTCDATETIME for consistent audit time. If this procedure is retained, change GETDATE to SYSUTCDATETIME and add error/row-count handling.")

    doc.add_page_break()
    doc.add_heading("Stored procedure recommendation", level=2)
    add_paragraph(doc,
                  "Do not move the explainable risk engine into a large stored procedure. Keep business decisions in AmlWorkflowService where they are versionable and testable. Stored procedures are most useful later for stable reporting queries, bulk watchlist import staging, audit export and operational maintenance.")
    for item in [
        "usp_AmlRecords_BulkImportFromStaging - validate and merge approved list versions.",
        "usp_ComplianceDashboard_GetSummary - aggregate high-volume dashboard metrics.",
        "usp_AuditLogs_GetPage - stable server-side audit paging for large history.",
        "usp_TransactionMonitoring_GetVelocityFacts - optional optimized feature query, while keeping final decision logic in C#.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("9. Data access, audit, soft delete and concurrency", level=1)
    access_rows = [
        ("AmlRecordRepository", "EF Core", "Paging, ownership scope, create/edit/soft delete, concurrency and transactional audit."),
        ("AmlWorkflowService", "EF Core", "Customer, transaction, screening, risk, decision, alert and case workflow."),
        ("ComplianceRuleRepository", "EF Core", "Client-scoped rule CRUD, soft delete and audit."),
        ("UserRepository", "Dapper inline SQL", "Login user, permission resolution, users, roles and per-user access."),
        ("RoleRepository", "Dapper inline SQL", "Role listing, creation and permission assignment."),
        ("SubscriptionRepository", "Dapper inline SQL", "Plan paging, CRUD, soft delete, row-version checks and audit."),
        ("MenuService", "Dapper inline SQL", "Loads active database menu items and filters them against user claims."),
        ("AuditService", "Dapper inline SQL", "Writes and reads client-scoped audit events."),
    ]
    add_table(doc, ["Component", "Technology", "Responsibility"], access_rows, [2500, 1900, 4960], font_size=8.8)

    doc.add_heading("Transaction boundaries", level=2)
    add_paragraph(doc,
                  "Create, update and delete operations that also write audit evidence use database transactions in the EF Core repositories/services. Either both the business change and audit event commit, or the operation rolls back. The login audit is deliberately best-effort so an audit storage problem does not reject valid credentials.")

    doc.add_heading("Soft delete", level=2)
    add_code_block(doc, """
DeleteFg   = 'Y'
DeleteDate = current UTC date/time
DeleteBy   = logged-in username
IsActive   = false where the entity has an active flag
""")
    add_paragraph(doc,
                  "Soft deletion preserves evidence. EF Core query filters hide deleted mapped records. SQL/Dapper queries must explicitly include DeleteFg = 'N' where applicable.")

    doc.add_heading("Optimistic concurrency", level=2)
    add_paragraph(doc,
                  "RowVersion changes automatically whenever SQL Server updates a row. Edit forms return the original RowVersion. If another user has updated the record first, the update is rejected and the user receives a conflict message instead of silently overwriting newer data.")

    doc.add_heading("10. Permissions and access-control matrix", level=1)
    matrix_rows = [
        ("Dashboard", "Authenticated", "View summary"),
        ("AML Records", "AMLRecords.View", "Create/Edit/Delete and ViewClient/ViewAll are separate"),
        ("Users", "Users.View", "Create requires Users.Create; access editing requires Roles.Manage"),
        ("Roles & Permissions", "Roles.Manage", "Create roles and assign permissions"),
        ("Subscription", "Subscription.View", "Create/Edit/Delete are separate"),
        ("Customers", "Customers.View", "Create/Edit/Delete are separate"),
        ("Transactions", "Transactions.View", "Create and Evaluate are separate"),
        ("Threshold Rules", "ComplianceRules.View", "Create/Edit/Delete are separate"),
        ("Alerts", "Alerts.View", "Assignment/resolution requires Alerts.Manage"),
        ("Cases", "Cases.View", "Investigation updates require Cases.Manage"),
        ("Audit Logs", "AuditLogs.View", "Read-only audit access"),
    ]
    add_table(doc, ["Menu", "Visibility permission", "Additional control"], matrix_rows, [2500, 2900, 3960], font_size=8.8)

    doc.add_page_break()
    doc.add_heading("Recommended role model", level=2)
    role_rows = [
        ("Platform Admin", "All permissions including AMLRecords.ViewAll; restrict to trusted platform operators."),
        ("Client Admin", "Users, roles, client-wide AML records, customer/transaction/rule/alert/case management, but no cross-client permission."),
        ("AML Analyst", "Client-wide AML view, customer/transaction evaluation, alerts and cases; no user/role administration."),
        ("Data Maintainer", "AML list create/edit and optional delete; no case or user administration."),
        ("Viewer/Auditor", "Read-only menus required for review; no create/edit/delete/manage permissions."),
    ]
    add_table(doc, ["Role", "Recommended access"], role_rows, [2500, 6860], font_size=9.2)

    doc.add_heading("11. Configuration, migration, UAT and production guidance", level=1)
    doc.add_heading("Configuration", level=2)
    config_rows = [
        ("ConnectionStrings:AMLDatabase", "SQL Server connection. Use a secret store outside Development."),
        ("AllowedHosts", "Use the real UAT or production hostname; never use wildcard in hosted environments."),
        ("DataProtection:ApplicationName", "Keeps authentication-cookie key usage scoped to this application."),
        ("DataProtection:KeysPath", "Persistent protected key directory required for hosted environments and multiple instances."),
        ("ASPNETCORE_ENVIRONMENT", "Development, UAT or Production selects the appropriate configuration."),
    ]
    add_table(doc, ["Setting", "Use"], config_rows, [3300, 6060], font_size=9.2)

    doc.add_heading("Database script order", level=2)
    for item in [
        "Back up the target database and verify the backup.",
        "Run ClientIsolation.sql when upgrading a database that lacks Clients/ClientId/OwnerUserId.",
        "Run AuditLogs.sql.",
        "Run RolesPermissions.sql.",
        "Run MenuItems.sql.",
        "Run PortalAmlWorkflow.sql to synchronize columns and create the workflow tables.",
        "Run schema and cross-client integration tests.",
        "Restart the application and sign in again so the authentication cookie contains current permissions.",
    ]:
        add_numbered(doc, item, migration_num_id)

    doc.add_heading("UAT checklist", level=2)
    for item in [
        "Use a separate UAT database and UAT hostname.",
        "Seed separate test clients, users, roles and representative AML records.",
        "Prove users cannot access another ClientId by changing URL IDs.",
        "Verify user permission grants and explicit denies after logout/login.",
        "Test ALLOW, REVIEW and BLOCK transaction scenarios.",
        "Verify alert resolution and case disposition requirements.",
        "Verify soft-deleted rows disappear but audit history remains.",
        "Review logs without exposing connection strings, hashes or sensitive KYC values.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("Production controls", level=2)
    for item in [
        "Use managed secrets, least-privilege database credentials and encrypted connections.",
        "Persist and back up data-protection keys; configure shared keys for multiple application instances.",
        "Centralize structured logs and health monitoring.",
        "Add database backup, restore testing, retention and disaster-recovery procedures.",
        "Review role assignments regularly and record privileged-access changes.",
        "Run security, load, regression and SQL integration tests in the release pipeline.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("12. Testing evidence, limitations and remaining work", level=1)
    test_rows = [
        ("Release build", "Passed", "0 errors and 0 warnings after workflow migration changes."),
        ("Integration tests", "Passed", "6 of 6 tests passed with SQL Server test connection enabled."),
        ("Login query", "Passed", "Admin user resolves with ClientId, active Admin role and 32 effective permissions."),
        ("Schema verification", "Passed", "Missing IsGranted/RowVersion columns added and seven workflow tables created."),
        ("Browser negative login", "Passed", "Incorrect test password returns Invalid username or password instead of a database-unavailable message."),
        ("Correct-password browser login", "User verification", "The user's real password was not requested or entered during automated verification."),
    ]
    add_table(doc, ["Check", "Result", "Evidence"], test_rows, [2500, 1500, 5360], font_size=8.9)

    doc.add_heading("Known limitations", level=2)
    limitations = [
        "AML list data is currently client-owned. A centralized platform normally separates global authoritative lists from client-private lists.",
        "ListVersion is currently manual-v1; automated source/version ingestion is not implemented.",
        "Name matching is a basic normalized Levenshtein implementation and lacks aliases, scripts/transliteration, identifiers and entity ownership matching.",
        "Receiver screening does not use receiver date of birth, nationality or identifiers because those fields are not captured.",
        "Threshold-rule approval is not true maker-checker.",
        "New customer and workflow lists need server-side pagination for production scale.",
        "Subscription plans do not yet enforce API entitlements or quotas.",
        "Search Screening, Reports and Settings permissions are legacy entries without active menu/controller implementations.",
    ]
    for item in limitations:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("Recommended next project: client API", level=2)
    api_rows = [
        ("Customer screening", "POST /api/v1/customers/screen", "Screen customer/KYC identity against active list versions."),
        ("Transaction evaluation", "POST /api/v1/aml/evaluate", "Idempotent evaluation returning ALLOW/REVIEW/BLOCK and evidence references."),
        ("Evaluation lookup", "GET /api/v1/evaluations/{correlationId}", "Retrieve the persisted explainable result."),
        ("Case lookup", "GET /api/v1/cases/{caseReference}", "Controlled client access to case status."),
        ("Health", "GET /health", "Readiness and dependency checks without exposing sensitive details."),
    ]
    add_table(doc, ["Capability", "Suggested endpoint", "Use"], api_rows, [2200, 2800, 4360], font_size=8.9)
    add_paragraph(doc,
                  "The API should add client credentials, subscription/entitlement enforcement, idempotency keys, correlation middleware, API rate limits, request/decision logs, mTLS or OAuth2, health checks, telemetry and retry-safe processing. It should call shared application services rather than duplicate portal rules.")

    doc.add_heading("Implementation priority", level=2)
    priorities = [
        "Separate global watchlist master data from client-private data and implement versioned imports.",
        "Add maker-checker approval for rules and watchlist publication.",
        "Add pagination to customers, transactions, alerts and cases.",
        "Add OwnerUserId/resource authorization where users must not share client-wide operational records.",
        "Add country/product/channel risk masters and validated customer risk configuration.",
        "Build the separate API project with credentials, subscriptions, idempotency and integration tests.",
        "Add evidence attachments, case event history and STR/SAR workflow.",
        "Complete performance, security, recovery and production-readiness testing.",
    ]
    for item in priorities:
        add_numbered(doc, item, priority_num_id)

    doc.add_heading("Appendix A. Project structure", level=1)
    add_code_block(doc, """
AmlScreening.Domain/
    Common/                  Base entity and client ownership contracts
    Entities/                User, Client, AML record, subscription and audit entities
    Entities/Compliance/     Rules, customers, transactions, matches, risks, decisions, alerts, cases

AmlScreening.Application/
    Interface/               Service and repository contracts
    Models/                  Login, paging, menu, audit and workflow result models
    Security/                Claim names and centralized permission constants
    Services/                Authentication orchestration

AmlScreening.Infrastructure/
    Data/                    EF Core DbContext and query filters
    Persistence/             SQL connection factory
    Repositories/            AML, users, roles, subscriptions and rules
    Services/                Menu, audit and AML workflow engine

AmlScreening/
    Controllers/             MVC endpoints and permission policies
    ViewModels/              Form and page models
    Views/                   Razor pages and shared forms/layout
    Services/                CurrentUserContext
    Security/                Resource authorization handler
    wwwroot/css/             Portal visual styles

Database/                    Migration and deployment scripts
AmlScreening.IntegrationTests/  Automated repository/schema/isolation tests
AmlScreening.DatabaseMigrator/  Controlled SQL script runner
""")

    doc.add_heading("Appendix B. Database scripts", level=1)
    scripts = [
        ("AmlRecords.sql", "Base AML record table definition and seed-friendly structure."),
        ("AmlRecords.SoftDelete.sql", "Adds soft-delete support to existing AML records."),
        ("ClientIsolation.sql", "Adds Clients, ClientId, OwnerUserId and client-scoped indexes/foreign keys."),
        ("AuditLogs.sql", "Creates or synchronizes audit storage."),
        ("RolesPermissions.sql", "Creates permission tables and seeds application permissions and baseline roles."),
        ("MenuItems.sql", "Creates and seeds the active database-driven navigation tree."),
        ("ComplianceRules.sql", "Creates the compliance rules table for a new database."),
        ("PortalAmlWorkflow.sql", "Synchronizes required columns and creates Customers, Transactions, Matches, Risks, Decisions, Alerts and Cases."),
        ("PlatformHardening.sql", "Broader hardening/reference script; review carefully before use on an existing database."),
    ]
    add_table(doc, ["Script", "Purpose"], scripts, [3000, 6360], font_size=9.0)

    doc.add_heading("Appendix C. Status glossary", level=1)
    glossary = [
        ("NEW", "Transaction captured but not yet evaluated."),
        ("CLEAR", "No qualifying AML screening match."),
        ("POTENTIAL_MATCH", "Screening score between 75 and below 95."),
        ("CONFIRMED_MATCH", "Screening score of 95 or higher in the prototype."),
        ("ALLOW", "Transaction cleared by current screening/risk/rule logic."),
        ("REVIEW", "Transaction held for analyst investigation."),
        ("BLOCK", "Confirmed match or blocking rule action."),
        ("LOW / MEDIUM / HIGH / CRITICAL", "Risk bands for scores 0-29 / 30-59 / 60-79 / 80-100."),
        ("DeleteFg = Y", "Soft-deleted record retained for evidence but hidden from normal queries."),
        ("RowVersion", "SQL Server concurrency token used to detect stale edits."),
    ]
    add_table(doc, ["Term", "Meaning"], glossary, [3000, 6360], font_size=9.2)

    doc.add_heading("Conclusion", level=1)
    add_paragraph(doc,
                  "The portal now provides a coherent multi-client AML operations prototype: controlled login, database-driven permissions and menus, watchlist records, KYC customers, transactions, explainable risk and decisions, alerts, cases and audit history. The next major boundary is the client API and production-grade data governance. Future development should preserve the current client isolation, permission policies, audit evidence and explainability rather than bypassing them in new controllers, SQL procedures or API endpoints.")

    # Core properties and save.
    props = doc.core_properties
    props.title = "AML Screening Portal Project Documentation"
    props.subject = "Architecture, menus, database, stored procedures and workflow"
    props.author = "AML Screening Project Team"
    props.keywords = "AML, compliance, portal, SQL Server, permissions, workflow"
    props.comments = "Generated from the implemented project and live development database catalog."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
